import asyncio

import pytest
from langchain_core.messages import AIMessageChunk, HumanMessage
from pydantic import ValidationError

from app.ai.agents import ProposedAction
from app.ai.ingestion import extract
from app.ai.providers import ModelProvider
from app.core.config import settings


def test_provider_stream_preserves_usage_and_propagates_failures():
    class ControlledModel:
        async def astream(self, messages):
            yield AIMessageChunk(content="Grounded ")
            yield AIMessageChunk(
                content="answer [S1]",
                usage_metadata={"input_tokens": 4, "output_tokens": 3, "total_tokens": 7},
            )

    provider = object.__new__(ModelProvider)
    provider.model, provider.usage = ControlledModel(), {}

    async def collect():
        return "".join([delta async for delta in provider.stream([HumanMessage(content="question")])])

    assert asyncio.run(collect()) == "Grounded answer [S1]"
    assert provider.usage["total_tokens"] == 7

    class BrokenModel:
        async def astream(self, messages):
            yield AIMessageChunk(content="Partial")
            raise RuntimeError("Provider disconnected")

    provider.model = BrokenModel()
    with pytest.raises(RuntimeError, match="disconnected"):
        asyncio.run(collect())


def test_structured_output_is_validated(monkeypatch):
    class InvalidStructuredModel:
        def with_structured_output(self, schema, method):
            return self

        async def ainvoke(self, messages):
            return {"title": "", "description": "bad", "source_labels": []}

    monkeypatch.setattr(settings(), "model_supports_structured", True)
    provider = object.__new__(ModelProvider)
    provider.model = InvalidStructuredModel()
    with pytest.raises(ValidationError):
        asyncio.run(provider.structured([HumanMessage(content="question")], ProposedAction))


def test_scanned_or_blank_pdf_requires_ocr(tmp_path):
    from pypdf import PdfWriter

    path = tmp_path / "blank.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.write(path)
    with pytest.raises(ValueError, match="OCR"):
        extract(path, ".pdf")


def test_docx_heading_and_markdown_structure(tmp_path):
    from docx import Document

    path = tmp_path / "policy.docx"
    document = Document()
    document.add_heading("Access control", 1)
    document.add_paragraph("MFA is required.")
    document.save(path)
    sections = extract(path, ".docx")
    assert sections[0].heading == "Access control"
    assert "MFA is required" in sections[0].text
    markdown = tmp_path / "policy.md"
    markdown.write_text("# Policy\nFirst rule\n## Exceptions\nNo shared accounts", encoding="utf-8")
    sections = extract(markdown, ".md")
    assert [s.heading for s in sections] == ["Policy", "Exceptions"]
